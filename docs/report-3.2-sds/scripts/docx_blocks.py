"""DOCX building blocks that reproduce the Report 3.2 SDS template formatting.

The template uses direct formatting (no table styles): header rows #BDD7EE, alternating #FFFFFF / #DEEAF1
rows, #AAAAAA cell borders, coloured 1×1 call-out tables with a thick left border, Courier New code boxes,
bullet List Paragraph (numId 2). This module copies those exact properties so the output matches the template.
"""
import copy
import re
import struct
import unicodedata
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

PAGE_TW = 9026  # usable width in twips (template tblW, portrait)
PAGE_TW_LANDSCAPE = 14638
MONO_SUBST = {"👁": "*", "🔒": "#", "⚑": "F", "✦": "*", "◉": "o", "📘": "", "🕒": "", "💬": "", "🔍": "", "⏱": "",
              "🗒": "", "⚙": "", "✓": "v", "✕": "x", "▬": "=", "▭": "-", "●": "o", "»": ">", "↺": "", "🛡": "", "🧪": "", "🎤": ""}


def mono_lines(lines):
    """Courier-safe ASCII: drop emoji (non-monospace glyphs) and re-pad right borders of │…│ box lines."""
    out = []
    for line in lines:
        for k, v in MONO_SUBST.items():
            line = line.replace(k, v)
        out.append(unicodedata.normalize("NFC", line))
    rules = [len(l) for l in out if l[:1] in "┌└├"]
    width = max(rules) if rules else 0
    fixed = []
    for l in out:
        if width and l.startswith("│") and l.rstrip().endswith("│") and len(l.rstrip()) != width:
            inner = l.rstrip()[:-1].rstrip()
            if len(inner) < width - 1:
                l = inner.ljust(width - 1) + "│"
        fixed.append(l)
    return fixed
CALLOUTS = {
    "teal": ("D6EEF0", "1F6B75", "1F6B75"),
    "blue": ("DEEAF1", "2E75B6", "1F3864"),
    "green": ("E2EFDA", "1E5C1E", "1E5C1E"),
    "code": ("F4E4D4", "5C3317", "5C3317"),
    "purple": ("EBE2F5", "5C3292", "5C3292"),
}
_TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def _runs_xml(text, base_rpr="", color="000000", size=None):
    """Inline markup: **bold**, `code` (Courier New)."""
    out = []
    for part in _TOKEN.split(text):
        if not part:
            continue
        rpr = base_rpr
        if part.startswith("**") and part.endswith("**"):
            part, rpr = part[2:-2], rpr + "<w:b/><w:bCs/>"
        elif part.startswith("`") and part.endswith("`"):
            part, rpr = part[1:-1], rpr + '<w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:cs="Courier New"/>'
        sz = f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/>' if size else ""
        out.append(f'<w:r><w:rPr>{rpr}<w:color w:val="{color}"/>{sz}</w:rPr>'
                   f'<w:t xml:space="preserve">{escape(part)}</w:t></w:r>')
    return "".join(out)


def _png_size(path):
    with open(path, "rb") as fh:
        head = fh.read(24)
    return struct.unpack(">II", head[16:24])


class SdsDocument:
    def __init__(self, template_path):
        self.doc = Document(template_path)
        body = self.doc.element.body
        self.sect = body.find(qn("w:sectPr"))
        for child in list(body.iterchildren()):
            if child is not self.sect:
                body.remove(child)
        self.figure_no = 0
        self.orient = "portrait"

    @property
    def page_tw(self):
        return PAGE_TW_LANDSCAPE if self.orient == "landscape" else PAGE_TW

    @staticmethod
    def _orient(sect, orient):
        pg = sect.find(qn("w:pgSz"))
        if orient == "landscape":
            pg.set(qn("w:w"), "16838"); pg.set(qn("w:h"), "11906"); pg.set(qn("w:orient"), "landscape")
        else:
            pg.set(qn("w:w"), "11906"); pg.set(qn("w:h"), "16838")
            if pg.get(qn("w:orient")) is not None:
                del pg.attrib[qn("w:orient")]

    def orientation(self, orient):
        """Close the current section (keeping its orientation) and continue in `orient` on a new page."""
        if orient == self.orient:
            return
        sect = copy.deepcopy(self.sect)
        self._orient(sect, self.orient)
        p = parse_xml(f'<w:p {nsdecls("w")}><w:pPr/></w:p>')
        p.find(qn("w:pPr")).append(sect)
        self.sect.addprevious(p)
        self.orient = orient

    # ---------------------------------------------------------------- paragraphs
    def _p(self, xml):
        el = parse_xml(f'<w:p {nsdecls("w")}>{xml}</w:p>')
        self.sect.addprevious(el)
        return el

    def title(self, text, sub):
        self._p(f'<w:pPr><w:spacing w:after="80"/><w:jc w:val="center"/></w:pPr>'
                + _runs_xml(text, "<w:b/><w:bCs/>", "1F3864", 60))
        self._p(f'<w:pPr><w:spacing w:after="240"/><w:jc w:val="center"/></w:pPr>' + _runs_xml(sub, "", "595959", 26))

    def heading(self, text, level):
        self._p(f'<w:pPr><w:pStyle w:val="Heading{level}"/><w:keepNext/></w:pPr>' + _runs_xml(text))

    def para(self, text, italic=False, size=None, color="000000", align=None, before=55, after=55, bold=False):
        rpr = ("<w:i/><w:iCs/>" if italic else "") + ("<w:b/><w:bCs/>" if bold else "")
        jc = f'<w:jc w:val="{align}"/>' if align else ""
        self._p(f'<w:pPr><w:spacing w:before="{before}" w:after="{after}"/>{jc}</w:pPr>' + _runs_xml(text, rpr, color, size))

    def labelled(self, label, text):
        self._p('<w:pPr><w:spacing w:before="55" w:after="55"/></w:pPr>'
                + _runs_xml(f"**{label}**  {text}"))

    def bullet(self, text):
        self._p('<w:pPr><w:pStyle w:val="ListParagraph"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="2"/></w:numPr>'
                '<w:spacing w:before="35" w:after="35"/></w:pPr>' + _runs_xml(text))

    def spacer(self):
        self._p("")

    def page_break(self):
        self._p('<w:r><w:br w:type="page"/></w:r>')

    # ---------------------------------------------------------------- tables
    def _widths(self, weights):
        total, tw = sum(weights), self.page_tw
        ws = [int(tw * w / total) for w in weights]
        ws[-1] += tw - sum(ws)
        return ws

    @staticmethod
    def _cell(text, width, fill, header=False, bold=False, mono=False, border="AAAAAA", sz=6):
        borders = "".join(f'<w:{side} w:val="single" w:sz="{sz}" w:space="0" w:color="{border}"/>'
                          for side in ("top", "left", "bottom", "right"))
        paras = []
        for line in str(text).split("\n"):
            if header:
                paras.append(f'<w:p><w:pPr><w:jc w:val="center"/></w:pPr>{_runs_xml(line, "<w:b/><w:bCs/>", "1F3864")}</w:p>')
            else:
                rpr = "<w:b/><w:bCs/>" if bold else ""
                if mono:
                    rpr += '<w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:cs="Courier New"/>'
                paras.append(f"<w:p>{_runs_xml(line, rpr, '000000', 17 if mono else None)}</w:p>")
        return (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/><w:tcBorders>{borders}</w:tcBorders>'
                f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/><w:tcMar><w:top w:w="80" w:type="dxa"/>'
                f'<w:left w:w="130" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:right w:w="130" w:type="dxa"/>'
                f"</w:tcMar></w:tcPr>{''.join(paras)}</w:tc>")

    def _table(self, widths, rows_xml):
        grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
        borders = "".join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
                          for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
        xml = (f'<w:tbl {nsdecls("w")}><w:tblPr><w:tblW w:w="{sum(widths)}" w:type="dxa"/><w:tblBorders>{borders}</w:tblBorders>'
               f'<w:tblLayout w:type="fixed"/><w:tblCellMar><w:left w:w="10" w:type="dxa"/><w:right w:w="10" w:type="dxa"/></w:tblCellMar>'
               f'<w:tblLook w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="0"/>'
               f"</w:tblPr><w:tblGrid>{grid}</w:tblGrid>{rows_xml}</w:tbl>")
        self.sect.addprevious(parse_xml(xml))
        self._p("")  # template keeps an empty Normal paragraph after every table

    def table(self, headers, rows, weights=None, mono_cols=()):
        weights = weights or [1] * len(headers)
        ws = self._widths(weights)
        head = ('<w:tr><w:trPr><w:tblHeader/><w:cantSplit/></w:trPr>'
                + "".join(self._cell(h, w, "BDD7EE", header=True) for h, w in zip(headers, ws)) + "</w:tr>")
        body = []
        for i, row in enumerate(rows):
            fill = "FFFFFF" if i % 2 == 0 else "DEEAF1"
            body.append('<w:tr><w:trPr><w:cantSplit/></w:trPr>'
                        + "".join(self._cell(c, w, fill, mono=j in mono_cols) for j, (c, w) in enumerate(zip(row, ws))) + "</w:tr>")
        self._table(ws, head + "".join(body))

    def kv(self, pairs, weights=(1651000, 4080510)):
        ws = self._widths(weights)
        rows = []
        for i, (k, v) in enumerate(pairs):
            fill = "DEEAF1" if i % 2 == 0 else "FFFFFF"
            rows.append('<w:tr><w:trPr><w:cantSplit/></w:trPr>' + self._cell(k, ws[0], fill, bold=True) + self._cell(v, ws[1], fill) + "</w:tr>")
        self._table(ws, "".join(rows))

    def callout(self, kind, title, lines, mono=False):
        fill, border, title_color = CALLOUTS[kind]
        tw = self.page_tw
        if mono:
            lines = mono_lines(lines)
        sides = "".join(f'<w:{s} w:val="single" w:sz="{18 if s == "left" else 8}" w:space="0" w:color="{border}"/>'
                        for s in ("top", "left", "bottom", "right"))
        paras = [f'<w:p><w:pPr><w:spacing w:after="50"/></w:pPr>{_runs_xml(title, "<w:b/><w:bCs/>", title_color)}</w:p>']
        for line in lines:
            if mono:
                paras.append('<w:p><w:pPr><w:spacing w:before="18" w:after="18"/></w:pPr>'
                             + _runs_xml(line.replace(" ", " "), '<w:rFonts w:ascii="Courier New" w:hAnsi="Courier New" w:cs="Courier New"/>', "1A1A1A", 16)
                             + "</w:p>")
            else:
                paras.append('<w:p><w:pPr><w:spacing w:before="22" w:after="22"/></w:pPr>'
                             + _runs_xml(line, "<w:i/><w:iCs/>", "404040", 17) + "</w:p>")
        cell = (f'<w:tc><w:tcPr><w:tcW w:w="{tw}" w:type="dxa"/><w:tcBorders>{sides}</w:tcBorders>'
                f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/><w:tcMar><w:top w:w="100" w:type="dxa"/>'
                f'<w:left w:w="180" w:type="dxa"/><w:bottom w:w="100" w:type="dxa"/><w:right w:w="180" w:type="dxa"/></w:tcMar>'
                f"</w:tcPr>{''.join(paras)}</w:tc>")
        self._table([tw], f"<w:tr>{cell}</w:tr>")

    # ---------------------------------------------------------------- figures
    def figure(self, png_path, caption, max_w=None, max_h=None):
        land = self.orient == "landscape"
        max_w = max_w or (10.1 if land else 6.27)
        max_h = max_h or (5.6 if land else 8.2)
        px_w, px_h = _png_size(png_path)
        w = max_w
        if px_h / px_w * w > max_h:
            w = max_h * px_w / px_h
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(png_path, width=Inches(w))
        self.figure_no += 1
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figure {self.figure_no} — {caption}")
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        return self.figure_no

    # ---------------------------------------------------------------- header / save
    def set_header(self, old, new):
        for section in self.doc.sections:
            for para in section.header.paragraphs:
                for run in para.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)

    def set_core(self, title, author, subject):
        cp = self.doc.core_properties
        cp.title, cp.author, cp.subject = title, author, subject

    def save(self, path):
        self._orient(self.sect, self.orient)
        self.doc.save(path)
